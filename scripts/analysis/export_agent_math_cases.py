#!/usr/bin/env python3
"""Export paired Agent Math evaluation cases to review-friendly XLSX/CSV files."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any, Iterable


DEFAULT_INPUT = Path(
    "evals/agent_math_results/repro_500_case_analysis/paired_results.jsonl"
)
DEFAULT_BASENAME = "agent_math_case_analysis"

CATEGORY_BOTH_PASS = "都对"
CATEGORY_AGENT_ONLY = "仅 Agent 对"
CATEGORY_FULL_SFT_ONLY = "仅 Full-SFT 对"
CATEGORY_BOTH_FAIL = "都错"
CATEGORY_ORDER = (
    CATEGORY_BOTH_PASS,
    CATEGORY_AGENT_ONLY,
    CATEGORY_FULL_SFT_ONLY,
    CATEGORY_BOTH_FAIL,
)

HEADERS = (
    "case_id",
    "分类",
    "expression",
    "prompt",
    "ground_truth",
    "full_sft_prediction",
    "agent_prediction",
    "full_sft_pass",
    "agent_pass",
    "full_sft_called_math",
    "agent_called_math",
    "full_sft_tool_call_count",
    "agent_tool_call_count",
    "available_tools",
    "full_sft_raw_response",
    "agent_raw_response",
    "full_sft_turns_and_tool_results",
    "agent_turns_and_tool_results",
)

ILLEGAL_XML_CHARS = re.compile(
    "[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x84\x86-\x9F]"
)
EXCEL_CELL_CHAR_LIMIT = 32767


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "将 eval_agent_math.py 生成的 paired_results.jsonl 导出为便于 case "
            "分析的 XLSX/CSV。"
        )
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_INPUT,
        help=f"paired_results.jsonl 路径，默认：{DEFAULT_INPUT}",
    )
    parser.add_argument(
        "--output_dir",
        type=Path,
        default=None,
        help="输出目录；默认与输入文件相同。",
    )
    parser.add_argument(
        "--basename",
        default=DEFAULT_BASENAME,
        help=f"输出文件名（不含扩展名），默认：{DEFAULT_BASENAME}",
    )
    parser.add_argument(
        "--format",
        choices=("xlsx", "csv", "both"),
        default="xlsx",
        help="输出格式，默认只生成带逐行颜色标注的 XLSX。",
    )
    parser.add_argument(
        "--full_sft_label",
        default="full_sft",
        help="paired_results.jsonl 中 Full-SFT 对应的 models 键。",
    )
    parser.add_argument(
        "--agent_label",
        default="agent",
        help="paired_results.jsonl 中 Agent 对应的 models 键。",
    )
    return parser.parse_args()


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        raise FileNotFoundError(f"找不到输入文件：{path}")

    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as file:
        for line_number, line in enumerate(file, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"第 {line_number} 行不是合法 JSON：{exc}") from exc
            if not isinstance(value, dict):
                raise ValueError(f"第 {line_number} 行必须是 JSON object。")
            rows.append(value)

    if not rows:
        raise ValueError(f"输入文件没有有效 case：{path}")
    return rows


def category_from_passes(full_sft_pass: bool, agent_pass: bool) -> str:
    if full_sft_pass and agent_pass:
        return CATEGORY_BOTH_PASS
    if agent_pass:
        return CATEGORY_AGENT_ONLY
    if full_sft_pass:
        return CATEGORY_FULL_SFT_ONLY
    return CATEGORY_BOTH_FAIL


def model_result(
    case: dict[str, Any], label: str, line_number: int
) -> dict[str, Any]:
    models = case.get("models")
    if not isinstance(models, dict):
        raise ValueError(f"第 {line_number} 个 case 缺少 models object。")
    result = models.get(label)
    if not isinstance(result, dict):
        available = ", ".join(str(key) for key in models) or "<empty>"
        raise ValueError(
            f"第 {line_number} 个 case 的 models 中没有 {label!r}；"
            f"当前键：{available}"
        )
    return result


def extract_available_tools(tools: Any) -> str:
    if not isinstance(tools, list):
        return ""
    names: list[str] = []
    for tool in tools:
        if not isinstance(tool, dict):
            continue
        function = tool.get("function")
        if isinstance(function, dict) and function.get("name"):
            names.append(str(function["name"]))
        elif tool.get("name"):
            names.append(str(tool["name"]))
    return ", ".join(names)


def format_json_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def format_turns(turns: Any) -> str:
    if not isinstance(turns, list):
        return format_json_value(turns)

    sections: list[str] = []
    for position, turn in enumerate(turns, start=1):
        if not isinstance(turn, dict):
            sections.append(f"Turn {position}:\n{format_json_value(turn)}")
            continue

        turn_number = turn.get("turn", position)
        lines = [f"Turn {turn_number}", "模型原始输出：", str(turn.get("raw_response", ""))]
        calls = turn.get("parsed_tool_calls")
        if isinstance(calls, list) and calls:
            lines.append("工具调用与返回：")
            for call_index, call in enumerate(calls, start=1):
                if not isinstance(call, dict):
                    lines.append(f"  {call_index}. {format_json_value(call)}")
                    continue
                name = call.get("name", "<unknown>")
                arguments = format_json_value(call.get("arguments"))
                tool_result = format_json_value(call.get("tool_result"))
                lines.append(f"  {call_index}. {name}({arguments}) -> {tool_result}")
        sections.append("\n".join(lines))
    return "\n\n".join(sections)


def count_tool_calls(result: dict[str, Any]) -> int:
    """Read both evaluator schemas: an integer count or a list of calls."""
    tool_calls = result.get("tool_calls")
    if isinstance(tool_calls, bool):
        return int(tool_calls)
    if isinstance(tool_calls, int):
        return max(tool_calls, 0)
    if isinstance(tool_calls, float) and tool_calls.is_integer():
        return max(int(tool_calls), 0)
    if isinstance(tool_calls, (list, tuple)):
        return len(tool_calls)
    if isinstance(tool_calls, dict):
        return 1
    if isinstance(tool_calls, str) and tool_calls.strip().isdigit():
        return int(tool_calls.strip())

    # Older/intermediate result files may omit the aggregate count. In that
    # case, reconstruct it from the per-turn parsed tool-call records.
    turns = result.get("turns")
    if not isinstance(turns, list):
        return 0
    count = 0
    for turn in turns:
        if not isinstance(turn, dict):
            continue
        parsed_calls = turn.get("parsed_tool_calls")
        if isinstance(parsed_calls, list):
            count += len(parsed_calls)
        elif isinstance(parsed_calls, dict):
            count += 1
    return count


def flatten_cases(
    cases: Iterable[dict[str, Any]], full_sft_label: str, agent_label: str
) -> list[dict[str, Any]]:
    flattened: list[dict[str, Any]] = []
    for line_number, case in enumerate(cases, start=1):
        full_sft = model_result(case, full_sft_label, line_number)
        agent = model_result(case, agent_label, line_number)
        full_sft_pass = bool(full_sft.get("passed", False))
        agent_pass = bool(agent.get("passed", False))

        flattened.append(
            {
                "case_id": case.get("case_id", f"case_{line_number:04d}"),
                "分类": category_from_passes(full_sft_pass, agent_pass),
                "expression": case.get("expression", ""),
                "prompt": case.get("prompt", ""),
                "ground_truth": case.get("ground_truth", ""),
                "full_sft_prediction": full_sft.get("prediction", ""),
                "agent_prediction": agent.get("prediction", ""),
                "full_sft_pass": full_sft_pass,
                "agent_pass": agent_pass,
                "full_sft_called_math": bool(full_sft.get("used_math_tool", False)),
                "agent_called_math": bool(agent.get("used_math_tool", False)),
                "full_sft_tool_call_count": count_tool_calls(full_sft),
                "agent_tool_call_count": count_tool_calls(agent),
                "available_tools": extract_available_tools(case.get("tools")),
                "full_sft_raw_response": full_sft.get("raw_response", ""),
                "agent_raw_response": agent.get("raw_response", ""),
                "full_sft_turns_and_tool_results": format_turns(full_sft.get("turns")),
                "agent_turns_and_tool_results": format_turns(agent.get("turns")),
            }
        )
    return flattened


def clean_text(value: Any, protect_formula: bool = False) -> Any:
    if not isinstance(value, str):
        return value
    value = ILLEGAL_XML_CHARS.sub("", value)
    if len(value) > EXCEL_CELL_CHAR_LIMIT:
        marker = "\n...[内容超过 Excel 单元格上限，已截断]"
        value = value[: EXCEL_CELL_CHAR_LIMIT - len(marker)] + marker
    if protect_formula and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=HEADERS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    header: clean_text(row.get(header, ""), protect_formula=True)
                    for header in HEADERS
                }
            )


def write_xlsx(path: Path, rows: list[dict[str, Any]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "生成 XLSX 需要 openpyxl。请先执行：\n"
            "python -m pip install -r scripts/analysis/requirements.txt"
        ) from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    cases_sheet = workbook.active
    cases_sheet.title = "Cases"

    navy = "1F4E78"
    white = "FFFFFF"
    text_color = "1F2937"
    border_color = "B8C4CE"
    category_fills = {
        CATEGORY_BOTH_PASS: "E2F0D9",
        CATEGORY_AGENT_ONLY: "DDEBF7",
        CATEGORY_FULL_SFT_ONLY: "FFF2CC",
        CATEGORY_BOTH_FAIL: "F4CCCC",
    }
    category_strong_fills = {
        CATEGORY_BOTH_PASS: "A9D18E",
        CATEGORY_AGENT_ONLY: "9DC3E6",
        CATEGORY_FULL_SFT_ONLY: "FFD966",
        CATEGORY_BOTH_FAIL: "E6B8B7",
    }
    thin_bottom = Border(bottom=Side(style="thin", color=border_color))

    # One row per case. Every populated data row receives its category color.
    for column_index, header in enumerate(HEADERS, start=1):
        cell = cases_sheet.cell(row=1, column=column_index, value=header)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_bottom

    for row_index, row in enumerate(rows, start=2):
        category = str(row["分类"])
        row_fill = PatternFill("solid", fgColor=category_fills[category])
        for column_index, header in enumerate(HEADERS, start=1):
            value = clean_text(row.get(header, ""), protect_formula=True)
            cell = cases_sheet.cell(row=row_index, column=column_index, value=value)
            cell.fill = row_fill
            cell.font = Font(color=text_color)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        status_cell = cases_sheet.cell(row=row_index, column=2)
        status_cell.fill = PatternFill("solid", fgColor=category_strong_fills[category])
        status_cell.font = Font(color=text_color, bold=True)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        cases_sheet.row_dimensions[row_index].height = 72

    last_row = len(rows) + 1
    last_column = get_column_letter(len(HEADERS))
    cases_sheet.auto_filter.ref = f"A1:{last_column}{last_row}"
    cases_sheet.freeze_panes = "C2"
    cases_sheet.sheet_view.showGridLines = False
    cases_sheet.row_dimensions[1].height = 30
    cases_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    cases_sheet.page_setup.fitToWidth = 1
    cases_sheet.page_setup.fitToHeight = 0

    widths = {
        "A": 15,
        "B": 14,
        "C": 30,
        "D": 48,
        "E": 16,
        "F": 19,
        "G": 17,
        "H": 15,
        "I": 12,
        "J": 21,
        "K": 18,
        "L": 22,
        "M": 20,
        "N": 22,
        "O": 58,
        "P": 58,
        "Q": 64,
        "R": 64,
    }
    for column, width in widths.items():
        cases_sheet.column_dimensions[column].width = width
    workbook.save(path)


def main() -> None:
    args = parse_args()
    input_path = args.input.expanduser().resolve()
    output_dir = (
        args.output_dir.expanduser().resolve()
        if args.output_dir is not None
        else input_path.parent
    )

    cases = load_jsonl(input_path)
    rows = flatten_cases(cases, args.full_sft_label, args.agent_label)
    counts = Counter(str(row["分类"]) for row in rows)

    generated: list[Path] = []
    if args.format in ("csv", "both"):
        csv_path = output_dir / f"{args.basename}.csv"
        write_csv(csv_path, rows)
        generated.append(csv_path)
    if args.format in ("xlsx", "both"):
        xlsx_path = output_dir / f"{args.basename}.xlsx"
        write_xlsx(xlsx_path, rows)
        generated.append(xlsx_path)

    print(f"已导出 {len(rows)} 个 case：", flush=True)
    for category in CATEGORY_ORDER:
        print(f"  {category}: {counts[category]}", flush=True)
    for path in generated:
        print(f"  {path}", flush=True)


if __name__ == "__main__":
    main()
